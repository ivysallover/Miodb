"""
Cómo lee el motor de Python un CSV (detect_csv_format + pd.read_csv, como `_read_dataframe`) y un
Excel (`read_excel_table`), para comparar contra el lector de Rust. Escribe los archivos de prueba en
`golden/csv/` y `golden/xlsx/` y guarda, por archivo, lo que detectó (codificación, separador, filas
salteadas) y la tabla que dejó pandas: nombres, tipo de cada columna y valores.

    python desktop/engine/tests/golden/make_golden_reader.py
"""
import datetime
import io
import json
import math
import sys
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "dashboard-ia" / "backend"))
import logging  # noqa: E402
logging.disable(logging.CRITICAL)
from services.analysis_pipeline import detect_csv_format, read_excel_table  # noqa: E402

HERE = Path(__file__).with_name("csv")
XLSX = Path(__file__).with_name("xlsx")

FILES = {
    "simple.csv": "fecha,ventas,cliente,unidades\n2024-01-01,1200.5,Ana,3\n2024-01-02,,Beto,\n2024-01-03,980,Carla,2\n2024-01-04,1100.25,,4\n",
    "punto_y_coma.csv": "Fecha;Importe;Región\n01/02/2024;1.234,56;Córdoba\n02/02/2024;980,10;Neuquén\n03/02/2024;12.000,00;Salta\n04/02/2024;5,5;Jujuy\n",
    "tabulado.tsv": "a\tb\tc\n1\t2\tx\n3\t4\ty\n5\t6\tz\n7\t8\tw\n",
    "barras.csv": "id|nombre|monto\n1|uno|10\n2|dos|20\n3|tres|30\n4|cuatro|40\n",
    "comillas.csv": 'nombre,comentario,valor\n"Pérez, Juan","dijo ""hola""",10\n"López","línea uno\nlínea dos",20\n"Gómez",simple,30\nSosa,"a,b,c",40\n',
    "vacios.csv": "a,b,c,d\nNA,1,null,x\nN/A,2,None,y\n,3,nan,-\nNULL,4,NaN,z\n#N/A,5,<NA>,w\n",
    "booleanos.csv": "a,b,c,d\nTrue,true,TRUE,True\nFalse,false,FALSE,\nTrue,true,FALSE,False\nFalse,false,TRUE,True\n",
    "enteros_con_huecos.csv": "id,cantidad,precio\n1,10,1.5\n2,,2.5\n3,30,\n4,40,4.0\n5,50,5\n",
    "ceros_adelante.csv": "codigo,cp,telefono\n00123,05000,0351-4567890\n00456,05001,0351-1112233\n00789,05002,0351-9998877\n01011,05003,0351-5554433\n",
    "cientifico.csv": "x,y,z\n1e3,inf,1.5\n2.5E-2,-inf,.5\n-3,1,5.\n4,2,-0.0\n",
    "titulo_arriba.csv": "Reporte de ventas 2024\nGenerado por el sistema\nfecha,ventas,region\n2024-01-01,10,N\n2024-01-02,20,S\n2024-01-03,30,E\n2024-01-04,40,O\n2024-01-05,50,N\n",
    "encabezados_repetidos.csv": "ventas,ventas,,ventas\n1,2,3,4\n5,6,7,8\n9,10,11,12\n",
    "filas_rotas.csv": "a,b,c\n1,2,3\n4,5\n6,7,8,9\n10,11,12\n13,14,15\n",
    "espacios.csv": "a,b,c\n 12 ,x , 1.5\n13, y,2.5 \n 14,z ,3.5\n15 , w,4.5\n",
    "crlf.csv": "a,b\r\n1,uno\r\n2,dos\r\n3,tres\r\n4,cuatro\r\n",
    "solo_encabezado.csv": "a,b,c\n",
    "una_columna.csv": "valor\n10\n20\n30\n40\n",
    "mezcla.csv": "a,b\n1,x\n2,3\nhola,4\n5,6\n",
    "lineas_en_blanco.csv": "a,b\n\n1,2\n\n3,4\n5,6\n\n",
    "titulo_punto_y_coma.csv": "Reporte de ventas\n\nfecha;ventas;region\n01/01/2024;10;N\n02/01/2024;20;S\n03/01/2024;30;E\n04/01/2024;40;O\n",
    "coma_final.csv": "a,b,c\n1,2,3,\n4,5,6,\n7,8,9,\n10,11,12,\n",
    "grandes.csv": "id,monto\n9223372036854775807,1\n123456789012345678901,2\n-5,3\n7,4\n",
}
LATIN1 = {
    "latin1.csv": ("fecha;descripción;importe\n01/03/2024;Camión de reparto;1500\n02/03/2024;Señal para año nuevo;980\n03/03/2024;Pequeño ajuste;120\n04/03/2024;Ñandú;45\n", "latin1"),
}
BOM = {"bom.csv": "﻿a,b\n1,2\n3,4\n5,6\n7,8\n"}
D = datetime.datetime
# Hojas de Excel: filas tal como se escriben (None = celda vacía), desde la celda de arriba a la izquierda.
EXCEL = {
    "tipos.xlsx": [
        ["entero", "decimal", "texto", "si_no", "fecha", "codigo", "vacios", "", "entero"],
        [1, 1.5, "Ana", True, D(2024, 1, 1), "00123", "N/A", 10, 7],
        [2, None, "Beto", False, D(2024, 1, 2), "00456", 3, 20, 8],
        [3, 2.25, None, True, D(2024, 1, 3, 10, 30), "00789", "NULL", 30, 9],
        [4, 3.0, "Dani", False, D(2024, 1, 4), "01011", 4, 40, 10],
    ],
    "ventas.xlsx": [["fecha", "ventas", "sucursal", "unidades"]]
    + [[D(2024, 1, 1) + datetime.timedelta(days=i), round(1000 + 37.5 * i + (i % 7) * 12.25, 2), "Norte" if i % 2 else "Sur", None if i == 5 else i % 9 + 1] for i in range(40)],
    "mezcla.xlsx": [["a", "b"], [1, "x"], [2, 3], ["hola", 4], [5, 6]],
    "titulo_arriba.xlsx": [["Reporte de ventas 2024"], [None], ["fecha", "ventas", None, "region"]]
    + [[D(2024, 2, i + 1), 10 * (i + 1), None, "NSEO"[i % 4]] for i in range(6)] + [[None], [None]],
}


def write_xlsx(path, rows, start_row=1, start_col=1):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=start_row + i, column=start_col + j, value=v)
    wb.create_sheet("otra").cell(row=1, column=1, value="no se lee")
    wb.save(path)


def value(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, float) and math.isinf(v):
        return "inf" if v > 0 else "-inf"
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.isoformat()
    if hasattr(v, "item"):
        v = v.item()
    return v


def main():
    HERE.mkdir(exist_ok=True)
    blobs = {name: text.encode("utf-8") for name, text in FILES.items()}
    blobs.update({name: text.encode(enc) for name, (text, enc) in LATIN1.items()})
    blobs.update({name: text.encode("utf-8") for name, text in BOM.items()})
    cases = []
    for name, data in blobs.items():
        (HERE / name).write_bytes(data)
        enc, sep, skip = detect_csv_format(io.BytesIO(data))
        try:
            df = pd.read_csv(io.BytesIO(data), encoding=enc, sep=sep, skiprows=skip, engine="c", on_bad_lines="skip")
            how = "c"
        except Exception as ex:  # noqa: BLE001
            try:
                df = pd.read_csv(io.BytesIO(data), encoding="latin1", sep=None, skiprows=skip, engine="python", on_bad_lines="skip")
                how = "python"
            except Exception as ex2:  # noqa: BLE001
                cases.append({"file": name, "encoding": enc, "sep": sep, "skiprows": skip, "error": f"{type(ex).__name__}/{type(ex2).__name__}"})
                print(f"{name:<28} ERROR {ex2}")
                continue
        cols = [{"name": str(c), "dtype": str(df[c].dtype), "values": [value(v) for v in df[c].tolist()]} for c in df.columns]
        cases.append({"file": name, "encoding": enc, "sep": sep, "skiprows": skip, "engine": how,
                      "index_is_default": isinstance(df.index, pd.RangeIndex), "rows": len(df), "columns": cols})
        print(f"{name:<28} enc={enc:<12} sep={sep!r:<5} skip={skip} filas={len(df)} tipos={[c['dtype'] for c in cols]}")
    XLSX.mkdir(exist_ok=True)
    corridas = {"corrida.xlsx": (EXCEL["mezcla.xlsx"], 3, 2), "titulo_corrido.xlsx": (EXCEL["titulo_arriba.xlsx"], 2, 3)}
    for name, (rows, r0, c0) in {**{n: (r, 1, 1) for n, r in EXCEL.items()}, **corridas}.items():
        write_xlsx(XLSX / name, rows, r0, c0)
        df = read_excel_table(XLSX / name)
        cols = [{"name": str(c), "dtype": str(df[c].dtype), "values": [value(v) for v in df[c].tolist()]} for c in df.columns]
        cases.append({"file": name, "excel": True, "rows": len(df), "columns": cols})
        print(f"{name:<28} excel filas={len(df)} tipos={[c['dtype'] for c in cols]}")
    out = Path(__file__).with_name("reader.json")
    out.write_text(json.dumps({"origen": "detect_csv_format + pd.read_csv (pandas %s)" % pd.__version__, "casos": cases}, indent=1, ensure_ascii=False))
    print(f"{len(cases)} archivos -> {out}")


if __name__ == "__main__":
    main()
